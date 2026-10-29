"""Parseo y validación offline del Excel "DATOS DE AGUATERÍA ACTUAL".

No toca la base de datos. Lee el .xlsx, normaliza clientes, inmuebles y deudas
con las mismas reglas del reporte `reporte_deudas` y escribe en --salida:

  problemas.csv     todo lo que no se pudo resolver solo (bloqueante o aviso)
  revision_E.csv    celdas DEUDA (col. E) con un solo monto: completar 'valor' con 2 o 3
                    (mismo formato que decisiones.csv; se pasa también con --decisiones)
  adelantos.csv     notas de la col. V (pagos adelantados) para tratar a mano
  resumen.txt       totales de control por concepto y por mes
  normalizado.json  datos listos para generar el staging (solo sin bloqueantes)

Uso:
  python3 importar.py ARCHIVO.xlsx [--decisiones decisiones.csv [otro.csv ...]] [--tarifa 25000]
"""
import argparse
import csv
import json
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

FILA_ENCABEZADO = 2
PRIMERA_FILA = 3
COL_COD, COL_NOMBRE, COL_DOC, COL_CEL, COL_DEUDA = 1, 2, 3, 4, 5
COL_PRIMER_MES, COL_ULTIMO_MES = 6, 19  # F:S
COL_NOTA = 22  # V: notas de pagos adelantados

CONCEPTO_CONSUMO, CONCEPTO_CONEXION, CONCEPTO_DEUDA_ANTERIOR = 1, 2, 3

# Campos que el usuario puede fijar en decisiones.csv (cod_inmueble,campo,valor)
CAMPOS_DECISION = {
    'documento',      # cédula/RUC correcto cuando el Excel trae dos números
    'concepto_E',     # 2 (Conexión) o 3 (Deuda anterior) para un monto suelto en DEUDA
    'tarifa',         # tarifa mensual propia del inmueble si no es la estándar
    'exceso_mes',     # 'deuda_anterior': el exceso sobre la tarifa pasa a Deuda anterior
}

RE_MARCA_ESTADO = re.compile(r'\s*\(?\s*(USUARIO\s+CORTADO|TAPONADO)\s*\)?', re.I)
RE_CORTADA = re.compile(r'^CONEXI[ÓO]N\s+CORTADA$', re.I)
RE_RUC = re.compile(r'^\d{5,9}-\d$')


class Problema:
    def __init__(self, fila, cod, tipo, detalle, bloqueante):
        self.fila, self.cod, self.tipo = fila, cod, tipo
        self.detalle, self.bloqueante = detalle, bloqueante

    def fila_csv(self):
        return [self.fila, self.cod, self.tipo, 'BLOQUEANTE' if self.bloqueante else 'AVISO', self.detalle]


def sin_acentos(texto):
    return ''.join(c for c in unicodedata.normalize('NFD', texto) if unicodedata.category(c) != 'Mn')


def limpiar_nombre(nombre):
    """Recorta, colapsa espacios y quita las marcas de estado del nombre."""
    nombre = RE_MARCA_ESTADO.sub('', str(nombre))
    return re.sub(r'\s+', ' ', nombre).strip()


def clave_nombre(nombre):
    """Nombre para comparar: sin acentos, mayúsculas y sin número final ('ELADIA VERÓN 2')."""
    n = re.sub(r'\s+\d+$', '', sin_acentos(limpiar_nombre(nombre)).upper())
    return re.sub(r'\s+', ' ', n).strip()


def normalizar_documento(valor):
    """Devuelve (documento, tipo, aviso). tipo: CI | RUC | SIN_TIPO.

    RUC (con guion) y sin documento se guardan sin tipo_documento, igual que
    cliente_page.dart; sin documento es '0'. Dos números en una celda no se
    adivinan: queda '0' y un aviso para que el usuario lo decida.
    """
    if valor is None:
        return '0', 'SIN_TIPO', None
    if isinstance(valor, (int, float)):
        return str(int(valor)), 'CI', None
    texto = str(valor).strip()
    if texto in ('', '-'):
        return '0', 'SIN_TIPO', None
    compacto = texto.replace('.', '').replace(' ', '')
    if RE_RUC.match(compacto):
        return compacto, 'RUC', None
    if compacto.isdigit() and texto.count(' ') == 0:
        return compacto, 'CI', None
    return '0', 'SIN_TIPO', f'documento ambiguo "{texto}"'


def normalizar_celular(valor):
    """Devuelve (celular, avisos). Solo dígitos, 10 con 09 al inicio; sin guion."""
    if valor is None or str(valor).strip() == '':
        return '', []
    texto = str(int(valor)) if isinstance(valor, (int, float)) else str(valor)
    validos, invalidos = [], []
    for token in re.split(r'[\s/;,]+', texto.strip()):
        digitos = re.sub(r'\D', '', token)
        if not digitos:
            continue
        if len(digitos) == 9 and digitos[0] == '9':
            digitos = '0' + digitos
        if len(digitos) == 10 and digitos.startswith('09'):
            validos.append(digitos)
        else:
            invalidos.append(token)
    avisos = []
    if len(validos) > 1:
        avisos.append(f'varios celulares "{texto.strip()}": se usa el primero')
    if invalidos:
        avisos.append(f'celular no válido "{" ".join(invalidos)}"')
    return (validos[0] if validos else ''), avisos


def montos_en_texto(valor):
    """Montos numéricos de una celda de texto ('30000\\n450000', '50000 775000')."""
    return [int(m) for m in re.findall(r'\d+', str(valor).replace('.', ''))]


def leer_decisiones(ruta):
    decisiones = defaultdict(dict)
    if not ruta or not Path(ruta).exists():
        return decisiones
    with open(ruta, newline='', encoding='utf-8') as f:
        lineas = (l for l in f if l.strip() and not l.lstrip().startswith('#'))
        for r in csv.DictReader(lineas):
            campo = r['campo'].strip()
            if not r['valor'].strip():
                continue  # plantilla sin completar (p. ej. revision_E.csv)
            if campo not in CAMPOS_DECISION:
                raise ValueError(f'decisiones.csv: campo desconocido "{campo}"')
            decisiones[r['cod_inmueble'].strip()][campo] = r['valor'].strip()
    return decisiones


def meses_del_encabezado(ws):
    meses = []
    for c in range(COL_PRIMER_MES, COL_ULTIMO_MES + 1):
        v = ws.cell(FILA_ENCABEZADO, c).value
        if not isinstance(v, datetime):
            raise ValueError(f'Encabezado {get_column_letter(c)}{FILA_ENCABEZADO} no es una fecha: {v!r}')
        meses.append(v.strftime('%Y-%m'))
    return meses


def procesar(ws, decisiones, tarifa_estandar=25000):
    """Devuelve dict con inmuebles, clientes, deudas, problemas, revision_E, adelantos."""
    meses = meses_del_encabezado(ws)
    problemas, revision_e, adelantos = [], [], []
    inmuebles, deudas = [], []
    clientes = {}
    cods_vistos = {}
    por_doc = defaultdict(list)  # documento -> [(clave_nombre, cod, fila)]

    def problema(fila, cod, tipo, detalle, bloqueante=False):
        problemas.append(Problema(fila, cod, tipo, detalle, bloqueante))

    for fila in range(PRIMERA_FILA, ws.max_row + 1):
        crudo_cod = ws.cell(fila, COL_COD).value
        crudo_nombre = ws.cell(fila, COL_NOMBRE).value
        if crudo_cod is None and (crudo_nombre is None or not str(crudo_nombre).strip()):
            continue
        cod = str(crudo_cod).strip() if crudo_cod is not None else ''
        if not cod or crudo_nombre is None:
            problema(fila, cod, 'FILA_INCOMPLETA', 'falta el n° de inmueble o el nombre', True)
            continue
        if cod in cods_vistos:
            problema(fila, cod, 'COD_DUPLICADO',
                     f'ya usado en la fila {cods_vistos[cod]}: corregir en el Excel', True)
            continue
        cods_vistos[cod] = fila
        dec = decisiones.get(cod, {})

        nombre = limpiar_nombre(crudo_nombre)
        es_cortada = bool(RE_CORTADA.match(nombre))
        celdas = [ws.cell(fila, c).value for c in range(COL_PRIMER_MES, COL_ULTIMO_MES + 1)]
        textos = [str(x) for x in celdas if isinstance(x, str) and x.strip()]
        marca_corte = bool(RE_MARCA_ESTADO.search(str(crudo_nombre))) or any(
            re.search(r'CORTAD', t, re.I) for t in textos)
        estado = 'DESCONECTADO' if (es_cortada or marca_corte) else 'CONECTADO'

        nota = ws.cell(fila, COL_NOTA).value
        if nota is not None and str(nota).strip():
            adelantos.append([fila, cod, nombre, str(nota).strip()])

        inmueble = {'cod_inmueble': cod, 'fila': fila, 'estado': estado,
                    'tarifa': int(dec.get('tarifa', tarifa_estandar)), 'cliente': None}

        if es_cortada:
            inmuebles.append(inmueble)
            continue

        # ---- cliente
        documento, tipo_doc, aviso_doc = normalizar_documento(ws.cell(fila, COL_DOC).value)
        if 'documento' in dec:
            documento, tipo_doc, _ = normalizar_documento(dec['documento'])
            aviso_doc = None
        if aviso_doc:
            problema(fila, cod, 'DOC_AMBIGUO', f'{aviso_doc} ({nombre}): se carga sin documento hasta decidir')
        celular, avisos_cel = normalizar_celular(ws.cell(fila, COL_CEL).value)
        for a in avisos_cel:
            problema(fila, cod, 'CELULAR', f'{a} ({nombre})')
        if '/' in nombre:
            problema(fila, cod, 'TITULARES_MULTIPLES', f'"{nombre}" se carga como una sola razón social')

        clave = clave_nombre(nombre)
        if documento != '0':
            # misma cédula y mismo nombre = mismo cliente con varios inmuebles
            clave_cliente = f'{documento}|{clave}'
            por_doc[documento].append((clave, cod, fila, nombre))
        else:
            clave_cliente = f'0|{cod}'
        if clave_cliente not in clientes:
            clientes[clave_cliente] = {
                'clave': clave_cliente, 'razon_social': re.sub(r'\s+\d+$', '', nombre), 'documento': documento,
                'tipo_documento': tipo_doc, 'celular': celular, 'inmuebles': [],
            }
        clientes[clave_cliente]['inmuebles'].append(cod)
        inmueble['cliente'] = clave_cliente
        inmuebles.append(inmueble)

        # ---- deudas de consumo (F:S)
        tarifa = inmueble['tarifa']
        exceso_anterior = 0
        for i, valor in enumerate(celdas):
            if not isinstance(valor, (int, float)) or valor <= 0:
                continue
            valor = int(valor)
            if valor > tarifa:
                if i == 0 or dec.get('exceso_mes') == 'deuda_anterior':
                    # sep-2025 trae arrastre de deuda vieja: tarifa a Consumo, resto a Deuda anterior
                    exceso_anterior += valor - tarifa
                    valor = tarifa
                else:
                    problema(fila, cod, 'MES_MAYOR_TARIFA',
                             f'{meses[i]} = {valor} supera la tarifa {tarifa} ({nombre}). '
                             'Indicar tarifa o exceso_mes=deuda_anterior en decisiones.csv', True)
                    continue
            deudas.append({'cod_inmueble': cod, 'concepto': CONCEPTO_CONSUMO, 'ciclo': meses[i],
                           'saldo': valor, 'monto': tarifa})
        if estado == 'CONECTADO' and celdas[-1] is None:
            problema(fila, cod, 'ULTIMO_MES_VACIO', f'{meses[-1]} vacío ({nombre}): se toma como sin cargo')
        if not any(isinstance(x, (int, float)) for x in celdas) and estado == 'CONECTADO':
            problema(fila, cod, 'SIN_MESES', f'sin ningún mes cargado ({nombre}) y sin marca de corte')

        # ---- columna DEUDA (E): Deuda anterior y Conexión
        anterior, conexion = exceso_anterior, 0
        e = ws.cell(fila, COL_DEUDA).value
        if isinstance(e, (int, float)) and e > 0:
            concepto = dec.get('concepto_E')
            if concepto in (str(CONCEPTO_DEUDA_ANTERIOR), str(CONCEPTO_CONEXION)):
                if int(concepto) == CONCEPTO_DEUDA_ANTERIOR:
                    anterior += int(e)
                else:
                    conexion += int(e)
            else:
                revision_e.append([cod, 'concepto_E', '', fila, nombre, int(e)])
                problema(fila, cod, 'DEUDA_SIN_CONCEPTO',
                         f'monto suelto {int(e)} en DEUDA ({nombre}): indicar concepto_E 2 o 3', True)
        elif isinstance(e, str) and e.strip() and e.strip().upper() != 'PAGADO':
            montos = montos_en_texto(e)
            if len(montos) == 2:
                anterior += montos[0]
                conexion += montos[1]
            else:
                problema(fila, cod, 'DEUDA_ILEGIBLE', f'celda DEUDA "{e.strip()}" ({nombre})', True)
        if anterior:
            deudas.append({'cod_inmueble': cod, 'concepto': CONCEPTO_DEUDA_ANTERIOR, 'ciclo': None,
                           'saldo': anterior, 'monto': anterior})
        if conexion:
            deudas.append({'cod_inmueble': cod, 'concepto': CONCEPTO_CONEXION, 'ciclo': None,
                           'saldo': conexion, 'monto': conexion})

    # cédulas repetidas con nombres distintos: no se fusionan, se anotan
    for documento, items in por_doc.items():
        if len({c for c, *_ in items}) > 1:
            detalle = '; '.join(f'fila {f} cód. {c} {n}' for _, c, f, n in items)
            problema(items[0][2], items[0][1], 'DOC_REPETIDO_NOMBRES_DISTINTOS',
                     f'documento {documento}: {detalle}')

    return {'meses': meses, 'inmuebles': inmuebles, 'clientes': list(clientes.values()),
            'deudas': deudas, 'problemas': problemas, 'revision_E': revision_e, 'adelantos': adelantos}


def totales(resultado):
    por_concepto, por_mes = defaultdict(int), defaultdict(int)
    for d in resultado['deudas']:
        por_concepto[d['concepto']] += d['saldo']
        if d['ciclo']:
            por_mes[d['ciclo']] += d['saldo']
    return por_concepto, por_mes


def escribir_salida(resultado, salida):
    salida = Path(salida)
    salida.mkdir(parents=True, exist_ok=True)

    def csv_(nombre, cabecera, filas):
        with open(salida / nombre, 'w', newline='', encoding='utf-8') as f:
            w = csv.writer(f)
            w.writerow(cabecera)
            w.writerows(filas)

    csv_('problemas.csv', ['fila', 'cod_inmueble', 'tipo', 'severidad', 'detalle'],
         [p.fila_csv() for p in resultado['problemas']])
    # mismo formato que decisiones.csv: completar 'valor' (2 Conexión, 3 Deuda anterior)
    csv_('revision_E.csv', ['cod_inmueble', 'campo', 'valor', 'fila', 'nombre', 'monto'],
         resultado['revision_E'])
    csv_('adelantos.csv', ['fila', 'cod_inmueble', 'nombre', 'nota'], resultado['adelantos'])

    por_concepto, por_mes = totales(resultado)
    bloqueantes = sum(p.bloqueante for p in resultado['problemas'])
    lineas = [
        f'Inmuebles: {len(resultado["inmuebles"])} '
        f'(conectados {sum(i["estado"] == "CONECTADO" for i in resultado["inmuebles"])}, '
        f'desconectados {sum(i["estado"] == "DESCONECTADO" for i in resultado["inmuebles"])})',
        f'Clientes: {len(resultado["clientes"])}',
        f'Deudas: {len(resultado["deudas"])}',
        f'Problemas: {bloqueantes} bloqueantes, {len(resultado["problemas"]) - bloqueantes} avisos',
        '',
        'Saldo por concepto:',
        f'  Consumo (1):        {por_concepto[1]:>12,}',
        f'  Conexión (2):       {por_concepto[2]:>12,}',
        f'  Deuda anterior (3): {por_concepto[3]:>12,}',
        f'  TOTAL:              {sum(por_concepto.values()):>12,}',
        '',
        'Consumo por mes:',
    ] + [f'  {m}: {por_mes[m]:>12,}' for m in resultado['meses']]
    (salida / 'resumen.txt').write_text('\n'.join(lineas) + '\n', encoding='utf-8')

    normalizado = salida / 'normalizado.json'
    if bloqueantes:
        normalizado.unlink(missing_ok=True)
    else:
        datos = {k: resultado[k] for k in ('meses', 'inmuebles', 'clientes', 'deudas')}
        normalizado.write_text(json.dumps(datos, ensure_ascii=False, indent=1), encoding='utf-8')
    return bloqueantes


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('xlsx')
    ap.add_argument('--decisiones', nargs='+', default=[str(Path(__file__).with_name('decisiones.csv'))])
    ap.add_argument('--salida', default=str(Path(__file__).with_name('out')))
    ap.add_argument('--tarifa', type=int, default=25000, help='tarifa mensual estándar (default 25000)')
    args = ap.parse_args(argv)

    ws = openpyxl.load_workbook(args.xlsx, data_only=True).active
    decisiones = defaultdict(dict)
    for ruta in args.decisiones:
        for cod, campos in leer_decisiones(ruta).items():
            decisiones[cod].update(campos)
    resultado = procesar(ws, decisiones, args.tarifa)
    bloqueantes = escribir_salida(resultado, args.salida)
    print((Path(args.salida) / 'resumen.txt').read_text(encoding='utf-8'))
    if bloqueantes:
        print(f'{bloqueantes} problemas bloqueantes: ver {args.salida}/problemas.csv. '
              'No se generó normalizado.json.')
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
