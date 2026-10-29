import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl

import importar as imp


def libro(filas):
    """Hoja con el encabezado real (sep-2025 a oct-2026) y las filas dadas.

    Cada fila: [cod, nombre, doc, cel, deuda, *14 meses] (se completa con None).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(14):
        ws.cell(2, 6 + i, datetime(2025 + (8 + i) // 12, (8 + i) % 12 + 1, 1))
    for r, fila in enumerate(filas, start=3):
        for c, v in enumerate(fila, start=1):
            ws.cell(r, c, v)
    return ws


def meses(**por_indice):
    m = [0] * 14
    for k, v in por_indice.items():
        m[int(k[1:])] = v
    return m


def correr(filas, decisiones=None):
    return imp.procesar(libro(filas), decisiones or {})


def tipos(res):
    return {p.tipo for p in res['problemas']}


class Normalizacion(unittest.TestCase):
    def test_documento(self):
        self.assertEqual(imp.normalizar_documento(2853288), ('2853288', 'CI', None))
        self.assertEqual(imp.normalizar_documento('4125450-3'), ('4125450-3', 'RUC', None))
        self.assertEqual(imp.normalizar_documento('4.237.043-4'), ('4237043-4', 'RUC', None))
        self.assertEqual(imp.normalizar_documento(None), ('0', 'SIN_TIPO', None))
        self.assertEqual(imp.normalizar_documento('-'), ('0', 'SIN_TIPO', None))
        doc, tipo, aviso = imp.normalizar_documento('5806326 / 7115644')
        self.assertEqual((doc, tipo), ('0', 'SIN_TIPO'))
        self.assertIn('ambiguo', aviso)
        self.assertIsNotNone(imp.normalizar_documento('4829098 6911047')[2])

    def test_celular(self):
        self.assertEqual(imp.normalizar_celular('0985-582258'), ('0985582258', []))
        self.assertEqual(imp.normalizar_celular(981311902), ('0981311902', []))
        self.assertEqual(imp.normalizar_celular('0971107-379')[0], '0971107379')
        cel, avisos = imp.normalizar_celular('0991-433349 0981-264959')
        self.assertEqual(cel, '0991433349')
        self.assertEqual(len(avisos), 1)
        cel, avisos = imp.normalizar_celular(8514565)
        self.assertEqual(cel, '')
        self.assertEqual(len(avisos), 1)
        self.assertEqual(imp.normalizar_celular(None), ('', []))

    def test_nombre(self):
        self.assertEqual(imp.limpiar_nombre('  DANIEL LÓPEZ (TAPONADO) '), 'DANIEL LÓPEZ')
        self.assertEqual(imp.limpiar_nombre('SADY BLANCO USUARIO CORTADO'), 'SADY BLANCO')
        self.assertEqual(imp.clave_nombre('ELADIA VERÓN 2'), imp.clave_nombre('Eladia  Verón'))

    def test_montos_en_texto(self):
        self.assertEqual(imp.montos_en_texto('30000\n450000'), [30000, 450000])
        self.assertEqual(imp.montos_en_texto('50000 775000'), [50000, 775000])


class Reglas(unittest.TestCase):
    def test_codigo_duplicado_bloquea(self):
        res = correr([[302, 'A', 1, '0981111111', None, *meses()],
                      [302, 'B', 2, '0981111112', None, *meses()]])
        dup = [p for p in res['problemas'] if p.tipo == 'COD_DUPLICADO']
        self.assertEqual(len(dup), 1)
        self.assertTrue(dup[0].bloqueante)
        self.assertEqual(len(res['inmuebles']), 1)

    def test_cortada_sin_cliente_ni_deuda(self):
        res = correr([[10, ' CONEXIÓN CORTADA']])
        self.assertEqual(res['inmuebles'][0]['estado'], 'DESCONECTADO')
        self.assertIsNone(res['inmuebles'][0]['cliente'])
        self.assertEqual((res['clientes'], res['deudas']), ([], []))

    def test_usuario_cortado_es_desconectado(self):
        fila = [3, 'CARMEN SOTO', 2853288, '0985-582258', 500000, 'USUARIO CORTADO'] + [None] * 13
        res = correr([fila], {'3': {'concepto_E': '3'}})
        self.assertEqual(res['inmuebles'][0]['estado'], 'DESCONECTADO')
        self.assertEqual([(d['concepto'], d['saldo']) for d in res['deudas']], [(3, 500000)])

    def test_consumo_por_mes_con_saldo(self):
        res = correr([[4, 'MARICELA', 1, '0994476316', None, *meses(i11=25000, i12=25000, i13=20000)]])
        consumo = [(d['ciclo'], d['saldo']) for d in res['deudas']]
        self.assertEqual(consumo, [('2026-08', 25000), ('2026-09', 25000), ('2026-10', 20000)])
        self.assertTrue(all(d['monto'] == 25000 for d in res['deudas']))

    def test_septiembre_grande_va_a_deuda_anterior(self):
        res = correr([[5, 'RUBÉN', 1, '0985426867', None, *meses(i0=900000)]])
        por = {(d['concepto'], d['ciclo']): d['saldo'] for d in res['deudas']}
        self.assertEqual(por, {(1, '2025-09'): 25000, (3, None): 875000})
        self.assertFalse(any(p.bloqueante for p in res['problemas']))

    def test_mes_mayor_a_tarifa_fuera_de_septiembre_bloquea(self):
        fila = [6, 'NELSY', 1, '0985426867', None, *meses(i13=50000)]
        res = correr([fila])
        self.assertIn('MES_MAYOR_TARIFA', tipos(res))
        res = correr([fila], {'6': {'tarifa': '50000'}})
        self.assertNotIn('MES_MAYOR_TARIFA', tipos(res))
        self.assertEqual(res['deudas'][0]['monto'], 50000)
        res = correr([fila], {'6': {'exceso_mes': 'deuda_anterior'}})
        por = {(d['concepto'], d['ciclo']): d['saldo'] for d in res['deudas']}
        self.assertEqual(por, {(1, '2026-10'): 25000, (3, None): 25000})

    def test_deuda_con_dos_montos_anterior_y_conexion(self):
        res = correr([[7, 'JORGELINA', 1, '0985426867', '30000\n450000', *meses()]])
        por = {d['concepto']: d['saldo'] for d in res['deudas']}
        self.assertEqual(por, {3: 30000, 2: 450000})

    def test_deuda_un_monto_pide_concepto(self):
        fila = [8, 'JOSÉ', 1, '0985426867', 600000, *meses()]
        res = correr([fila])
        self.assertIn('DEUDA_SIN_CONCEPTO', tipos(res))
        self.assertEqual(res['revision_E'][0][:3], ['8', 'concepto_E', ''])
        res = correr([fila], {'8': {'concepto_E': '2'}})
        self.assertEqual([(d['concepto'], d['saldo']) for d in res['deudas']], [(2, 600000)])

    def test_deuda_pagado_no_genera_deuda(self):
        res = correr([[9, 'MIRIAN', 1, '0985426867', 'PAGADO', *meses()]])
        self.assertEqual(res['deudas'], [])
        self.assertFalse(any(p.bloqueante for p in res['problemas']))

    def test_mismo_documento_y_nombre_un_solo_cliente(self):
        res = correr([[20, 'ELADIA VERÓN 2', 6205285, '0986320427', None, *meses()],
                      [22, 'ELADIA VERÓN', 6205285, '0986320427', None, *meses()]])
        self.assertEqual(len(res['clientes']), 1)
        self.assertEqual(res['clientes'][0]['razon_social'], 'ELADIA VERÓN')
        self.assertEqual(res['clientes'][0]['inmuebles'], ['20', '22'])

    def test_mismo_documento_nombres_distintos_no_fusiona_y_avisa(self):
        res = correr([[188, 'LIZ ALICE GARCETE RAMOS', 7208444, None, None, *meses()],
                      [201, 'ALICE GARCETE RAMOS', 7208444, None, None, *meses()]])
        self.assertEqual(len(res['clientes']), 2)
        self.assertIn('DOC_REPETIDO_NOMBRES_DISTINTOS', tipos(res))
        self.assertFalse(any(p.bloqueante for p in res['problemas']))

    def test_sin_documento_un_cliente_por_fila(self):
        res = correr([[1, 'A', None, None, None, *meses()], [2, 'B', None, None, None, *meses()]])
        self.assertEqual(len(res['clientes']), 2)
        self.assertEqual({c['documento'] for c in res['clientes']}, {'0'})

    def test_documento_ambiguo_se_resuelve_con_decision(self):
        fila = [125, 'RICHARD / RUMILDA', '5806326 / 7115644', None, None, *meses()]
        res = correr([fila])
        self.assertIn('DOC_AMBIGUO', tipos(res))
        res = correr([fila], {'125': {'documento': '5806326'}})
        self.assertNotIn('DOC_AMBIGUO', tipos(res))
        self.assertEqual(res['clientes'][0]['documento'], '5806326')

    def test_nota_de_adelanto_se_reporta(self):
        wb = libro([[41, 'OFICINA', None, None, None, *meses()]])
        wb.cell(3, imp.COL_NOTA, 'NOVIEMBRE DICIEMBRE 2026')
        res = imp.procesar(wb, {})
        self.assertEqual(res['adelantos'][0][1], '41')

    def test_salida_sin_bloqueantes_genera_json_y_con_bloqueantes_no(self):
        with tempfile.TemporaryDirectory() as tmp:
            ok = correr([[1, 'A', 1, '0985426867', None, *meses(i13=25000)]])
            self.assertEqual(imp.escribir_salida(ok, tmp), 0)
            self.assertTrue((Path(tmp) / 'normalizado.json').exists())
            self.assertIn('25,000', (Path(tmp) / 'resumen.txt').read_text(encoding='utf-8'))
            mal = correr([[1, 'A', 1, '0985426867', 100, *meses()]])
            self.assertEqual(imp.escribir_salida(mal, tmp), 1)
            self.assertFalse((Path(tmp) / 'normalizado.json').exists())


if __name__ == '__main__':
    unittest.main()
